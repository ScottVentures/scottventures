"""Local processing implementations for ScottPDF's public tools."""

from __future__ import annotations

import difflib
import html
import ipaddress
import json
import mimetypes
import os
import socket
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse
from urllib.request import HTTPRedirectHandler, Request, build_opener

import fitz
from bs4 import BeautifulSoup
from docx import Document as WordDocument
from docx.shared import Inches
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches as PptInches
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from PIL import Image
from io import BytesIO


MAX_PAGES = int(os.getenv("SCOTTPDF_MAX_PAGES", "1000"))


class ToolError(Exception):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


@dataclass
class Result:
    path: Path
    filename: str
    media_type: str


SUPPORTED_TOOLS = {
    "merge", "split", "compress", "jpg_to_pdf", "word_to_pdf", "powerpoint_to_pdf",
    "excel_to_pdf", "html_to_pdf", "pdf_to_jpg", "pdf_to_word", "pdf_to_powerpoint",
    "pdf_to_excel", "convert_pdf_to_pdfa", "remove_pages", "extract", "organize_pdf",
    "scan_pdf", "repair_pdf", "ocr_pdf", "rotate_pdf", "add_pdf_page_number",
    "pdf_add_watermark", "crop_pdf", "edit_pdf", "pdf_forms", "unlock_pdf",
    "protect_pdf", "sign_pdf", "redact_pdf", "compare_pdf", "pdf_summarize",
    "translate_pdf", "pdf_to_markdown",
}


def _base(value: str, fallback: str) -> str:
    value = Path(value or "").name
    value = re.sub(r"[^A-Za-z0-9._ -]", "", value).strip(" .")
    return Path(value).stem[:80] or fallback


def _range_groups(spec: str, count: int) -> list[list[int]]:
    if not spec.strip():
        return [[index] for index in range(count)]
    groups: list[list[int]] = []
    for item in spec.split(","):
        match = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d+)\s*)?", item)
        if not match:
            raise ToolError("Enter page numbers or ranges such as 1-3,5,8-9.")
        first = int(match.group(1))
        last = int(match.group(2) or first)
        if first < 1 or last < first or last > count:
            raise ToolError(f"Page ranges must be between 1 and {count}.")
        groups.append(list(range(first - 1, last)))
    return groups


def _ranges(spec: str, count: int) -> list[int]:
    return [index for group in _range_groups(spec, count) for index in group]


def _public_web_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise ToolError("Enter a public http or https address.")
    if parsed.port and parsed.port not in {80, 443}:
        raise ToolError("Only standard website ports are allowed.")
    try:
        addresses = {item[4][0] for item in socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))}
    except OSError as exc:
        raise ToolError("The website could not be reached.", 422) from exc
    if not addresses or any(not ipaddress.ip_address(address).is_global for address in addresses):
        raise ToolError("Private network addresses cannot be converted.", 422)
    return url


class _NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


def _fetch_html(url: str) -> str:
    opener = build_opener(_NoRedirect())
    current = _public_web_url(url)
    for _ in range(5):
        try:
            response = opener.open(Request(current, headers={"User-Agent": "ScottPDF/1.0", "Accept": "text/html"}), timeout=15)
        except Exception as exc:
            location = getattr(exc, "headers", {}).get("Location") if getattr(exc, "headers", None) else None
            if location:
                current = _public_web_url(urljoin(current, location))
                continue
            raise ToolError("The website could not be loaded.", 422) from exc
        with response:
            if "text/html" not in response.headers.get("Content-Type", "").lower():
                raise ToolError("The address did not return an HTML page.", 415)
            raw = response.read(3 * 1024 * 1024 + 1)
            if len(raw) > 3 * 1024 * 1024:
                raise ToolError("The HTML page is larger than the 3 MB limit.", 413)
            return raw.decode(response.headers.get_content_charset() or "utf-8", errors="replace")
    raise ToolError("The website redirected too many times.", 422)


def _html_to_pdf(source: str, output: Path) -> Result:
    soup = BeautifulSoup(source, "html.parser")
    for tag in soup(["script", "style", "noscript", "iframe", "object", "embed"]):
        tag.decompose()
    for image in soup.find_all("img"):
        image.decompose()
    story = fitz.Story(
        html=str(soup),
        user_css="body { font-family: sans-serif; font-size: 11pt; color: #222; } img { display: none; }",
    )
    media = fitz.paper_rect("a4")
    bounds = media + (36, 36, -36, -36)
    writer = fitz.DocumentWriter(str(output))
    more = True
    while more:
        device = writer.begin_page(media)
        more, _ = story.place(bounds)
        story.draw(device)
        writer.end_page()
    writer.close()
    return Result(output, output.name, "application/pdf")


def _convert_with_libreoffice(source: Path, work: Path) -> Path | None:
    executable = os.getenv("SCOTTPDF_LIBREOFFICE", "") or shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        return None
    try:
        subprocess.run(
            [executable, "--headless", "--convert-to", "pdf", "--outdir", str(work), str(source)],
            check=True, capture_output=True, timeout=90,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    converted = work / f"{source.stem}.pdf"
    return converted if converted.is_file() else None


def _pdf(path: Path, password: str = "") -> fitz.Document:
    try:
        document = fitz.open(path)
        if document.is_encrypted:
            if not password or not document.authenticate(password):
                raise ToolError("This PDF is password-protected. Enter its password to continue.", 422)
        if not document.is_pdf:
            raise ToolError("The uploaded file is not a PDF.", 415)
        if len(document) == 0:
            raise ToolError("The PDF has no pages.", 422)
        if len(document) > MAX_PAGES:
            raise ToolError(f"A PDF may contain at most {MAX_PAGES} pages.", 413)
        return document
    except ToolError:
        raise
    except Exception as exc:
        raise ToolError("ScottPDF could not read this PDF. It may be damaged or use an unsupported feature.", 422) from exc


def _save_pdf(document: fitz.Document, output: Path, **kwargs: Any) -> Result:
    try:
        document.save(output, garbage=4, deflate=True, **kwargs)
    except Exception as exc:
        raise ToolError("ScottPDF could not write the output PDF.", 422) from exc
    return Result(output, output.name, "application/pdf")


def _zip_result(files: list[tuple[Path, str]], output: Path) -> Result:
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for source, name in files:
            archive.write(source, name)
    return Result(output, output.name, "application/zip")


def _text_pdf(text_pages: list[str], output: Path, page_size=A4) -> Result:
    styles = getSampleStyleSheet()
    normal = styles["BodyText"]
    normal.leading = 14
    story: list[Any] = []
    for index, text in enumerate(text_pages):
        if index:
            story.append(PageBreak())
        for paragraph in text.split("\n"):
            if paragraph.strip():
                story.append(Paragraph(html.escape(paragraph), normal))
                story.append(Spacer(1, 3 * mm))
    SimpleDocTemplate(str(output), pagesize=page_size).build(story or [Paragraph(" ", normal)])
    return Result(output, output.name, "application/pdf")


def _pdf_text(document: fitz.Document) -> list[str]:
    return [page.get_text("text", sort=True) for page in document]


def _extract_text(documents: list[fitz.Document]) -> str:
    return "\n\n".join("\n".join(_pdf_text(doc)) for doc in documents)


def _docx_from_pdf(document: fitz.Document, output: Path) -> Result:
    word = WordDocument()
    for page_index, page in enumerate(document):
        if page_index:
            word.add_page_break()
        text = page.get_text("text", sort=True)
        for line in text.splitlines():
            word.add_paragraph(line)
    word.save(output)
    return Result(output, output.name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def _pdf_to_pptx(document: fitz.Document, output: Path) -> Result:
    deck = Presentation()
    for page in document:
        image = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes("png")
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(__import__("io").BytesIO(image), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(output)
    return Result(output, output.name, "application/vnd.openxmlformats-officedocument.presentationml.presentation")


def _pdf_to_xlsx(document: fitz.Document, output: Path) -> Result:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for page_index, page in enumerate(document, start=1):
        sheet = workbook.create_sheet(f"Page {page_index}"[:31])
        for row_index, line in enumerate(page.get_text("text", sort=True).splitlines(), start=1):
            cells = [cell.strip() for cell in re.split(r"\s{2,}|\t", line) if cell.strip()]
            if cells:
                for col_index, value in enumerate(cells, start=1):
                    sheet.cell(row_index, col_index, value[:32000])
    if not workbook.sheetnames:
        workbook.create_sheet("Page 1")
    workbook.save(output)
    return Result(output, output.name, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _images_to_pdf(paths: list[Path], output: Path, orientation: str = "auto") -> Result:
    document = fitz.open()
    for path in paths:
        try:
            with Image.open(path) as source:
                image = source.convert("RGB")
                buffer = BytesIO()
                image.save(buffer, format="JPEG", quality=94, optimize=True)
                image_width, image_height = image.size
        except Exception as exc:
            raise ToolError(f"Could not read image {path.name}. Use JPG, PNG, or another supported image.", 415) from exc

        if orientation == "portrait":
            page_width, page_height = fitz.paper_size("a4")
        elif orientation == "landscape":
            page_height, page_width = fitz.paper_size("a4")
        else:
            page_width, page_height = image_width * 72 / 150, image_height * 72 / 150
        page = document.new_page(width=page_width, height=page_height)
        margin = min(28, page_width * 0.06, page_height * 0.06)
        available = fitz.Rect(margin, margin, page_width - margin, page_height - margin)
        scale = min(available.width / image_width, available.height / image_height)
        draw_width, draw_height = image_width * scale, image_height * scale
        left = available.x0 + (available.width - draw_width) / 2
        top = available.y0 + (available.height - draw_height) / 2
        draw_rect = fitz.Rect(left, top, left + draw_width, top + draw_height)
        page.insert_image(draw_rect, stream=buffer.getvalue())
    return _save_pdf(document, output)


def _paragraph_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _summary(text: str, target_words: int = 180) -> str:
    sentences = re.split(r"(?<=[.!?])\s+", text)
    sentences = [s.strip() for s in sentences if len(s.strip()) > 25]
    if not sentences:
        return text[:4000] or "No extractable text was found in this PDF."
    words = re.findall(r"[A-Za-z0-9]{3,}", text.lower())
    frequency: dict[str, int] = {}
    for word in words:
        frequency[word] = frequency.get(word, 0) + 1
    ranked = sorted(
        enumerate(sentences),
        key=lambda pair: sum(frequency.get(w, 0) for w in re.findall(r"[A-Za-z0-9]{3,}", pair[1].lower())) / max(1, len(pair[1].split())),
        reverse=True,
    )
    chosen: list[tuple[int, str]] = []
    total = 0
    for position, sentence in ranked:
        if total + len(sentence.split()) > target_words and chosen:
            continue
        chosen.append((position, sentence))
        total += len(sentence.split())
        if total >= target_words:
            break
    return "\n\n".join(sentence for _, sentence in sorted(chosen))


def _ai_chat(prompt: str, system: str) -> str:
    api_key = os.getenv("SCOTTPDF_AI_API_KEY", "").strip()
    if not api_key:
        raise ToolError("AI tools need SCOTTPDF_AI_API_KEY configured on the server.", 503)
    import httpx

    base_url = os.getenv("SCOTTPDF_AI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    model = os.getenv("SCOTTPDF_AI_MODEL", "gpt-4.1-mini")
    try:
        response = httpx.post(
            f"{base_url}/chat/completions",
            headers={"Authorization": f"Bearer {api_key}"},
            json={"model": model, "messages": [{"role": "system", "content": system}, {"role": "user", "content": prompt}]},
            timeout=120,
        )
        response.raise_for_status()
        return response.json()["choices"][0]["message"]["content"].strip()
    except Exception as exc:
        raise ToolError("The configured AI service could not process this request.", 502) from exc


def process_tool(tool: str, paths: list[Path], names: list[str], options: dict[str, Any], work: Path) -> Result:
    if tool not in SUPPORTED_TOOLS:
        raise ToolError("This PDF tool is not implemented yet.", 404)
    stem = _base(str(options.get("output_name", "")), f"scottpdf-{tool.replace('_', '-')}")
    output = work / f"{stem}.pdf"

    if tool in {"jpg_to_pdf", "scan_pdf"}:
        if not paths:
            raise ToolError("Choose at least one image file.")
        return _images_to_pdf(paths, output, str(options.get("orientation", "auto")))

    if tool in {"word_to_pdf", "powerpoint_to_pdf", "excel_to_pdf"}:
        source = paths[0]
        converted = _convert_with_libreoffice(source, work)
        if converted:
            if converted != output:
                shutil.copyfile(converted, output)
            return Result(output, output.name, "application/pdf")
        if tool == "word_to_pdf":
            docx = WordDocument(source)
            pages = ["\n".join(p.text for p in docx.paragraphs)]
            for table in docx.tables:
                pages.append("\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows))
            return _text_pdf(pages, output)
        if tool == "powerpoint_to_pdf":
            presentation = Presentation(source)
            pages = []
            for slide in presentation.slides:
                lines = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
                pages.append("\n".join(lines))
            return _text_pdf(pages, output, landscape(A4))
        book = load_workbook(source, read_only=True, data_only=True)
        pages = []
        for sheet in book.worksheets:
            rows = ["\t".join(str(cell) if cell is not None else "" for cell in row) for row in sheet.iter_rows(values_only=True)]
            pages.append(f"{sheet.title}\n" + "\n".join(rows))
        return _text_pdf(pages, output, landscape(A4))

    if tool == "html_to_pdf":
        source = paths[0].read_text(encoding="utf-8", errors="replace") if paths else _fetch_html(str(options.get("url", "")))
        return _html_to_pdf(source, output)

    if tool == "compare_pdf":
        if len(paths) != 2:
            raise ToolError("Choose exactly two PDFs to compare.")
        left, right = (_pdf(path) for path in paths)
        diff = difflib.unified_diff(
            _extract_text([left]).splitlines(), _extract_text([right]).splitlines(),
            fromfile=names[0], tofile=names[1], lineterm="",
        )
        result = work / f"{stem}.txt"
        result.write_text("\n".join(diff) or "No text differences were found.", encoding="utf-8")
        return Result(result, result.name, "text/plain; charset=utf-8")

    if tool == "pdf_summarize":
        doc = _pdf(paths[0])
        content = _extract_text([doc])[:90000]
        target_words = int(options.get("target_words", 180))
        if target_words not in {90, 180, 300}:
            target_words = 180
        summary = _ai_chat(content, "Summarize the document accurately, using concise bullets and preserving important names, dates, and conclusions.") if os.getenv("SCOTTPDF_AI_API_KEY") else _summary(content, target_words)
        result = work / f"{stem}.txt"
        result.write_text(summary, encoding="utf-8")
        return Result(result, result.name, "text/plain; charset=utf-8")

    if tool == "translate_pdf":
        doc = _pdf(paths[0])
        source_text = _extract_text([doc])[:90000]
        target = str(options.get("target_language", "English"))[:60]
        translated = _ai_chat(source_text, f"Translate the supplied document into {target}. Preserve its meaning, headings, lists, and page breaks where possible. Return only the translation.")
        result = work / f"{stem}.txt"
        result.write_text(translated, encoding="utf-8")
        return Result(result, result.name, "text/plain; charset=utf-8")

    if tool == "pdf_to_markdown":
        doc = _pdf(paths[0])
        parts = []
        for index, page in enumerate(doc, start=1):
            parts.extend([f"## Page {index}", "", page.get_text("text", sort=True).strip(), ""])
        result = work / f"{stem}.md"
        result.write_text("\n".join(parts), encoding="utf-8")
        return Result(result, result.name, "text/markdown; charset=utf-8")

    if tool == "pdf_to_jpg":
        doc = _pdf(paths[0])
        exported = []
        image_format = str(options.get("format", "jpg")).lower()
        if image_format not in {"jpg", "png"}:
            image_format = "jpg"
        for index, page in enumerate(doc, start=1):
            image = work / f"page-{index:03d}.{image_format}"
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            pix.save(image, jpg_quality=90) if image_format == "jpg" else pix.save(image)
            exported.append((image, image.name))
        return _zip_result(exported, work / f"{stem}-images.zip")

    if tool in {"pdf_to_word", "pdf_to_powerpoint", "pdf_to_excel"}:
        doc = _pdf(paths[0])
        if tool == "pdf_to_word":
            return _docx_from_pdf(doc, work / f"{stem}.docx")
        if tool == "pdf_to_powerpoint":
            return _pdf_to_pptx(doc, work / f"{stem}.pptx")
        return _pdf_to_xlsx(doc, work / f"{stem}.xlsx")

    if tool == "convert_pdf_to_pdfa":
        raise ToolError("PDF/A conversion is not available in this build. It needs a dedicated archival validation engine.", 501)

    if tool == "ocr_pdf":
        doc = _pdf(paths[0])
        language = re.sub(r"[^a-zA-Z+_-]", "", str(options.get("language", "eng"))) or "eng"
        tessdata = os.getenv("TESSDATA_PREFIX")
        tesseract = os.getenv("TESSERACT_CMD", "tesseract")
        if not (Path(tesseract).is_file() if os.path.dirname(tesseract) else shutil.which(tesseract)):
            raise ToolError("OCR needs Tesseract and its language data installed on the server. See README-backend.md.", 503)
        searchable = fitz.open()
        try:
            for page in doc:
                pix = page.get_pixmap(dpi=200, alpha=False)
                converted = fitz.open("pdf", pix.pdfocr_tobytes(language=language, tessdata=tessdata))
                searchable.insert_pdf(converted)
                converted.close()
        except Exception as exc:
            raise ToolError("OCR failed. Check that Tesseract and the selected language pack are installed.", 503) from exc
        return _save_pdf(searchable, output)

    if tool == "jpg_to_pdf" or tool == "scan_pdf":
        raise ToolError("Choose image files for this tool.")

    if tool == "unlock_pdf":
        doc = _pdf(paths[0], str(options.get("password", "")))
        return _save_pdf(doc, output, encryption=fitz.PDF_ENCRYPT_NONE)

    if tool == "protect_pdf":
        password = str(options.get("password", ""))
        if len(password) < 4:
            raise ToolError("Choose a password with at least 4 characters.")
        doc = _pdf(paths[0])
        return _save_pdf(doc, output, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password)

    if tool == "merge":
        if len(paths) < 2:
            raise ToolError("Choose at least two PDFs to merge.")
        doc = fitz.open()
        for path in paths:
            doc.insert_pdf(_pdf(path))
        return _save_pdf(doc, output)

    if tool == "split":
        doc = _pdf(paths[0])
        groups = _range_groups(str(options.get("pages", "")), len(doc))
        if len(groups) == 1:
            part = fitz.open()
            part.insert_pdf(doc, from_page=groups[0][0], to_page=groups[0][-1])
            return _save_pdf(part, output)
        archive_files = []
        for count, group in enumerate(groups, start=1):
            part = fitz.open()
            for index in group:
                part.insert_pdf(doc, from_page=index, to_page=index)
            path = work / f"part-{count}.pdf"
            part.save(path, garbage=4, deflate=True)
            archive_files.append((path, f"{stem}-{count}.pdf"))
        return _zip_result(archive_files, work / f"{stem}.zip")

    doc = _pdf(paths[0], str(options.get("password", "")))

    if tool == "compress":
        level = str(options.get("level", "recommended"))
        image_options = {
            "extreme": (1100, 55),
            "recommended": (1600, 72),
            "less": (2400, 86),
        }
        max_dimension, quality = image_options.get(level, image_options["recommended"])
        seen_images: set[int] = set()
        for page in doc:
            for image_info in page.get_images(full=True):
                xref = image_info[0]
                if xref <= 0 or xref in seen_images:
                    continue
                seen_images.add(xref)
                try:
                    image_data = doc.extract_image(xref)["image"]
                    with Image.open(BytesIO(image_data)) as image:
                        image = image.convert("RGB")
                        if max(image.size) > max_dimension:
                            scale = max_dimension / max(image.size)
                            image = image.resize((max(1, int(image.width * scale)), max(1, int(image.height * scale))), Image.Resampling.LANCZOS)
                        encoded = BytesIO()
                        image.save(encoded, format="JPEG", quality=quality, optimize=True)
                    for image_page in doc:
                        if any(info[0] == xref for info in image_page.get_images(full=True)):
                            image_page.replace_image(xref, stream=encoded.getvalue())
                except Exception:
                    continue
        save_options = {"garbage": 4, "deflate": True, "deflate_images": True, "deflate_fonts": True, "use_objstms": 1}
        try:
            doc.save(output, **save_options)
        except Exception as exc:
            raise ToolError("ScottPDF could not compress this PDF.", 422) from exc
        if output.stat().st_size >= paths[0].stat().st_size:
            shutil.copyfile(paths[0], output)
        return Result(output, output.name, "application/pdf")

    if tool == "repair_pdf":
        try:
            doc.save(output, garbage=4, deflate=True, clean=True)
        except Exception as exc:
            raise ToolError("This PDF is too damaged for ScottPDF to repair.", 422) from exc
        return Result(output, output.name, "application/pdf")

    if tool in {"remove_pages", "extract", "organize_pdf", "rotate_pdf", "add_pdf_page_number", "pdf_add_watermark", "crop_pdf", "edit_pdf", "pdf_forms", "sign_pdf", "redact_pdf"}:
        if tool == "remove_pages":
            pages = _ranges(str(options.get("pages", "")), len(doc))
            if len(pages) >= len(doc):
                raise ToolError("A PDF must keep at least one page.")
            for index in sorted(set(pages), reverse=True):
                doc.delete_page(index)
        elif tool == "extract":
            pages = _ranges(str(options.get("pages", "")), len(doc))
            result_doc = fitz.open()
            for index in pages:
                result_doc.insert_pdf(doc, from_page=index, to_page=index)
            doc.close()
            doc = result_doc
        elif tool == "organize_pdf":
            raw = str(options.get("order", ""))
            order = [int(value.strip()) - 1 for value in raw.split(",") if value.strip()]
            if not order or len(order) != len(doc) or set(order) != set(range(len(doc))):
                raise ToolError(f"Enter every page once in the new order, for example 3,1,2. This PDF has {len(doc)} pages.")
            reordered = fitz.open()
            for index in order:
                reordered.insert_pdf(doc, from_page=index, to_page=index)
            doc.close()
            doc = reordered
        elif tool == "rotate_pdf":
            angle = int(options.get("angle", 90))
            if angle not in {90, 180, 270}:
                raise ToolError("Choose a rotation of 90, 180, or 270 degrees.")
            for index in _ranges(str(options.get("pages", "")), len(doc)):
                doc[index].set_rotation((doc[index].rotation + angle) % 360)
        elif tool == "add_pdf_page_number":
            start = max(1, int(options.get("start_number", 1)))
            position = str(options.get("position", "bottom-center"))
            for index, page in enumerate(doc):
                rect = page.rect
                point = (rect.width / 2, rect.height - 22) if position == "bottom-center" else (rect.width - 50, rect.height - 22) if position == "bottom-right" else (rect.width / 2, 25)
                page.insert_text(point, str(start + index), fontsize=10, fontname="helv", color=(0.2, 0.2, 0.2))
        elif tool == "pdf_add_watermark":
            text = str(options.get("text", "SCOTTPDF"))[:100]
            opacity = min(1.0, max(0.1, float(options.get("opacity", 0.25))))
            for page in doc:
                fontsize = min(52, page.rect.width / max(4, len(text) * 0.55))
                center = fitz.Point(page.rect.width / 2, page.rect.height / 2)
                text_width = fitz.Font("helv").text_length(text, fontsize=fontsize)
                baseline = fitz.Point(center.x - text_width / 2, center.y)
                page.insert_text(
                    baseline,
                    text,
                    fontsize=fontsize,
                    fontname="helv",
                    color=(0.5, 0.5, 0.5),
                    fill_opacity=opacity,
                    morph=(center, fitz.Matrix(-45)),
                )
        elif tool == "crop_pdf":
            top, right, bottom, left = (max(0, min(30, float(options.get(key, 0)))) / 100 for key in ("top", "right", "bottom", "left"))
            for page in doc:
                r = page.rect
                page.set_cropbox(fitz.Rect(r.x0 + r.width * left, r.y0 + r.height * top, r.x1 - r.width * right, r.y1 - r.height * bottom))
        elif tool in {"edit_pdf", "sign_pdf"}:
            text = str(options.get("text", ""))[:200]
            if not text:
                raise ToolError("Enter the text to add to the PDF.")
            page_index = max(0, min(len(doc) - 1, int(options.get("page", 1)) - 1))
            page = doc[page_index]
            x = max(0, min(page.rect.width - 20, float(options.get("x", 72))))
            y = max(20, min(page.rect.height - 10, float(options.get("y", page.rect.height / 2))))
            page.insert_text((x, y), text, fontsize=18 if tool == "sign_pdf" else 12, fontname="tiit" if tool == "sign_pdf" else "helv", color=(0.1, 0.1, 0.25))
        elif tool == "redact_pdf":
            terms = [value.strip() for value in str(options.get("terms", "")).splitlines() if value.strip()]
            if not terms:
                raise ToolError("Enter at least one word or phrase to permanently remove.")
            hits = 0
            for page in doc:
                for term in terms:
                    for rect in page.search_for(term, quads=False):
                        page.add_redact_annot(rect, fill=(0, 0, 0))
                        hits += 1
                if page.first_annot:
                    page.apply_redactions()
            if not hits:
                raise ToolError("None of the entered text was found. No redactions were applied.")
        elif tool == "pdf_forms":
            try:
                values = json.loads(str(options.get("fields", "{}")))
            except json.JSONDecodeError as exc:
                raise ToolError("Enter form values as JSON, for example {\"Name\":\"Scott\"}.") from exc
            found = 0
            for page in doc:
                for widget in page.widgets() or []:
                    if widget.field_name in values:
                        widget.field_value = str(values[widget.field_name])
                        widget.update()
                        found += 1
            if not found:
                raise ToolError("No matching PDF form fields were found.")
        return _save_pdf(doc, output)

    raise ToolError("This PDF tool is not implemented yet.", 501)
